from pathlib import Path
from docx import Document
from openpyxl import Workbook

EXPORT_PATH = Path(__file__).resolve().parents[2] / "data" / "exports"

# 导出md格式的文件
def export_markdown(content: str, filename: str) -> Path:
    EXPORT_PATH.mkdir(parents=True, exist_ok=True)

    file_path = EXPORT_PATH / f"{filename}.md"

    with open(file_path, "w", encoding="utf-8") as file:
        file.write(content)

    return file_path

# 导出word文档
def export_word(content: str, filename: str) -> Path:
    EXPORT_PATH.mkdir(parents=True, exist_ok=True)

    file_path = EXPORT_PATH / f"{filename}.docx"

    document = Document()

    for line in content.splitlines():
        line = line.strip()

        if not line:
            continue

        if line.startswith("# "):
            document.add_heading(line[2:], level=1)
        elif line.startswith("## "):
            document.add_heading(line[3:], level=2)
        elif line.startswith("### "):
            document.add_heading(line[4:], level=3)
        elif line.startswith("- "):
            document.add_paragraph(line[2:], style="List Bullet")
        else:
            document.add_paragraph(line)

    document.save(file_path)

    return file_path

# 导出Excel文件
def export_excel(
    destination: str,
    days: int,
    content: str,
    filename: str
) -> Path:
    EXPORT_PATH.mkdir(parents=True, exist_ok=True)

    file_path = EXPORT_PATH / f"{filename}.xlsx"

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "旅游行程"

    sheet.append([
        "目的地",
        "天数",
        "行程内容"
    ])

    sheet.append([
        destination,
        days,
        content
    ])

    sheet.column_dimensions["A"].width = 20
    sheet.column_dimensions["B"].width = 10
    sheet.column_dimensions["C"].width = 100

    workbook.save(file_path)

    return file_path